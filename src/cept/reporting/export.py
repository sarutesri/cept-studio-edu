"""Export a study report to PDF and DOCX.

Two different strategies, because the HTML report is interactive
(Plotly + ECharts) and neither PDF nor DOCX can hold that natively:

* **PDF** renders the *actual* ``report.html`` through a headless browser
  (Playwright/Chromium) and prints it — a faithful static snapshot of what
  the researcher sees on screen, SLD and charts included. Plotly.js/
  ECharts.js are vendored and inlined into the HTML (see
  ``reporting/render.py``), so this works fully offline.
* **DOCX** is built natively from the same `Case`/`StudyResult` data the
  HTML report uses (not scraped from the HTML) — a text-editable Word
  document with a static matplotlib rendering of the single-line diagram
  and result charts, since python-docx can't embed interactive JS.
"""

from __future__ import annotations

from pathlib import Path
from typing import Optional

from cept.schema.case import Case
from cept.schema.result import StudyResult


def export_pdf(html_path: str | Path, out_path: str | Path) -> Path:
    """Render an existing HTML report to PDF via headless Chromium."""
    from playwright.sync_api import sync_playwright

    html_path = Path(html_path).resolve()
    out_path = Path(out_path)
    if not html_path.exists():
        raise FileNotFoundError(f"HTML report not found: {html_path}")

    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page()
        page.goto(html_path.as_uri(), wait_until="networkidle")
        # 'networkidle' only means no more network requests — Plotly/ECharts
        # still need a beat to paint after their JS finishes running, or the
        # PDF captures a near-empty page. Verified: without this wait, a
        # full report with charts printed to an ~880-byte (empty) PDF.
        page.wait_for_timeout(1000)
        page.pdf(
            path=str(out_path),
            format="A4",
            print_background=True,
            margin={"top": "12mm", "bottom": "12mm", "left": "10mm", "right": "10mm"},
        )
        browser.close()
    return out_path


def _sld_figure(result: StudyResult):
    """Static matplotlib rendering of the (already engine-neutral)
    :class:`SLDModel` — works identically for OpenDSS or PowerFactory
    results since both populate the same node/edge schema."""
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    sld = result.sld
    fig, ax = plt.subplots(figsize=(7, 5))
    if sld is None or not sld.nodes:
        ax.text(0.5, 0.5, "No SLD data", ha="center", va="center")
        ax.axis("off")
        return fig

    pos = {n.id: (n.x, n.y) for n in sld.nodes}
    for edge in sld.edges:
        if edge.src in pos and edge.dst in pos:
            x1, y1 = pos[edge.src]
            x2, y2 = pos[edge.dst]
            style = "--" if edge.kind == "transformer" else "-"
            ax.plot([x1, x2], [y1, y2], style, color="#4a90d9", linewidth=1.5, zorder=1)

    for node in sld.nodes:
        x, y = pos[node.id]
        v = node.v_mean
        color = (
            "#2e7d32" if v is None else ("#c62828" if (v < sld.v_min_pu or v > sld.v_max_pu) else "#2e7d32")
        )
        ax.scatter([x], [y], s=180, color=color, edgecolor="black", zorder=2)
        label = node.id if v is None else f"{node.id}\n{v:.3f} pu"
        ax.annotate(label, (x, y), textcoords="offset points", xytext=(8, 6), fontsize=8)

    ax.set_title(sld.title)
    ax.axis("off")
    fig.tight_layout()
    return fig


def _result_chart_figure(result: StudyResult):
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    fig, ax = plt.subplots(figsize=(7, 3.5))
    if result.fault is not None:
        currents = result.fault.currents
        labels = [f"phase {c.phase}" for c in currents]
        values = [c.i_amp for c in currents]
        ax.bar(labels, values, color="#c62828")
        ax.set_ylabel("Fault current (A)")
        ax.set_title(f"Fault current @ {result.fault.bus} ({result.fault.fault_type})")
    elif result.load_flow is not None:
        agg: dict[str, list[float]] = {}
        for bv in result.load_flow.bus_voltages:
            agg.setdefault(bv.bus, []).append(bv.v_pu)
        buses = list(agg)
        means = [sum(v) / len(v) for v in agg.values()]
        ax.bar(buses, means, color="#2e7d32")
        ax.axhline(1.0, color="gray", linewidth=0.8, linestyle=":")
        ax.set_ylabel("Voltage (pu)")
        ax.set_title("Bus voltage profile")
        ax.tick_params(axis="x", rotation=60)
    else:
        ax.text(0.5, 0.5, "No chartable result data", ha="center", va="center")
        ax.axis("off")
    fig.tight_layout()
    return fig


def export_docx(
    result: StudyResult,
    case: Case,
    out_path: str | Path,
    *,
    validation_summary: Optional[dict] = None,
) -> Path:
    """Build a native, text-editable Word report from Case + StudyResult."""
    import matplotlib

    matplotlib.use("Agg")
    from docx import Document
    from docx.shared import Cm, Pt

    out_path = Path(out_path)
    doc = Document()

    doc.add_heading(case.meta.name, level=0)
    doc.add_paragraph(case.meta.description or "(no description provided)")

    doc.add_heading("Study parameters", level=1)
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    rows = [
        ("Study type", result.study_type),
        ("Engine", f"{result.engine} ({result.engine_version})"),
        ("Network kind", case.network.kind),
        ("Frequency", f"{case.network.frequency_hz} Hz"),
        ("Voltage limits", f"{case.standards.v_min_pu}-{case.standards.v_max_pu} pu"),
        ("Case fingerprint", case.fingerprint()),
        ("Generated", result.created_at),
    ]
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = str(value)

    if case.assumptions:
        doc.add_heading("Assumed values", level=1)
        warning = doc.add_paragraph(
            "Demonstrator inputs below were not supplied as research data and "
            "must not be treated as measured or project-specific values."
        )
        warning.runs[0].bold = True
        at = doc.add_table(rows=1, cols=3)
        at.style = "Light Grid Accent 1"
        hdr = at.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text = "Path", "Value", "Source"
        for assumption in case.assumptions:
            cells = at.add_row().cells
            cells[0].text = assumption.path
            cells[1].text = str(assumption.value)
            cells[2].text = assumption.source

    doc.add_heading("Single-line diagram", level=1)
    sld_fig = _sld_figure(result)
    sld_path = out_path.with_suffix(".sld.png")
    sld_fig.savefig(sld_path, dpi=150)
    doc.add_picture(str(sld_path), width=Cm(16))
    sld_path.unlink(missing_ok=True)
    import matplotlib.pyplot as plt

    plt.close(sld_fig)

    doc.add_heading("Results", level=1)
    chart_fig = _result_chart_figure(result)
    chart_path = out_path.with_suffix(".chart.png")
    chart_fig.savefig(chart_path, dpi=150)
    doc.add_picture(str(chart_path), width=Cm(16))
    chart_path.unlink(missing_ok=True)
    plt.close(chart_fig)

    if result.load_flow is not None:
        doc.add_heading("Bus voltages", level=2)
        vt = doc.add_table(rows=1, cols=4)
        vt.style = "Light Grid Accent 1"
        hdr = vt.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "Bus", "Phase", "V (pu)", "Angle (deg)"
        for bv in sorted(result.load_flow.bus_voltages, key=lambda b: (b.bus, b.phase)):
            cells = vt.add_row().cells
            cells[0].text = bv.bus
            cells[1].text = str(bv.phase)
            cells[2].text = f"{bv.v_pu:.4f}"
            cells[3].text = f"{bv.v_angle_deg:.2f}"

    if result.fault is not None:
        doc.add_heading("Fault currents", level=2)
        ft = doc.add_table(rows=1, cols=3)
        ft.style = "Light Grid Accent 1"
        hdr = ft.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text = "Phase", "Current (A)", "Angle (deg)"
        for c in result.fault.currents:
            cells = ft.add_row().cells
            cells[0].text = str(c.phase)
            cells[1].text = f"{c.i_amp:.2f}"
            cells[2].text = f"{c.i_angle_deg:.2f}"

    if validation_summary is not None:
        doc.add_heading("Validation", level=1)
        vt2 = doc.add_table(rows=1, cols=3)
        vt2.style = "Light Grid Accent 1"
        hdr = vt2.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text = "Check", "Value", "Status"
        for check in validation_summary.get("checks", []):
            cells = vt2.add_row().cells
            cells[0].text = check["label"]
            cells[1].text = str(check["value"])
            cells[2].text = "PASS" if check["passed"] else "WARNING"
        p = doc.add_paragraph()
        run = p.add_run(
            "This validation is an acceptance gate for the generated run, not a "
            "full engineering review. Do not treat missing project-specific "
            "topology or controller data as validated."
        )
        run.italic = True
        run.font.size = Pt(9)

    doc.save(str(out_path))
    return out_path


__all__ = ["export_pdf", "export_docx"]
